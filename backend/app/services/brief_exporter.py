"""DOCX and PDF exports of the designer brief."""

from __future__ import annotations

from pathlib import Path

from app.services.designer import DesignerBrief

ACTION_LABEL = {"keep": "KEEP", "change": "CHANGE", "add": "ADD", "remove": "REMOVE", "unknown": "UNKNOWN"}
ACTION_HEX = {"keep": "15803D", "change": "1D4ED8", "add": "7C3AED", "remove": "B91C1C", "unknown": "B45309"}


def _panel_name(p):
    return (p or "").replace("_", " ")


def export_docx(brief: DesignerBrief, path: Path) -> Path:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.7)
        s.top_margin = s.bottom_margin = Inches(0.6)
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    def shade(cell, hex_fill):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        cell._tc.get_or_add_tcPr().append(shd)

    def table(headers, rows, widths=None):
        t = d.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = h
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            shade(c, "1E293B")
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = "" if v is None else str(v)
                for p in cells[i].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        return t

    title = d.add_paragraph()
    run = title.add_run("Packaging Design Brief")
    run.font.size = Pt(20)
    run.font.bold = True
    sub = d.add_paragraph()
    sub.add_run(f"{brief.product}  ·  {brief.market}  ·  {brief.case}").font.size = Pt(11)
    meta = d.add_paragraph()
    r = meta.add_run(f"Status: {brief.status.upper()}   ·   Updated: {brief.updated_at[:16].replace('T', ' ')} UTC"
                     f"   ·   Readiness: {brief.readiness.score if brief.readiness else '-'}%")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(100, 116, 139)
    for line in brief.summary:
        d.add_paragraph(line, style="List Bullet")

    d.add_heading("1. What to do", level=1)
    t = table(["Action", "Item", "Panel", "Instruction", "Current → New"],
              [[ACTION_LABEL.get(i.action, i.action), i.item, _panel_name(i.panel), i.instruction,
                (f"{i.current or ''} → {i.new or ''}" if i.action == "change" else (i.new or i.current or ""))[:220]]
               for i in brief.items], [0.8, 1.4, 0.9, 2.2, 1.9])
    for row, item in zip(t.rows[1:], brief.items):
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(ACTION_HEX.get(item.action, "334155"))

    d.add_heading("2. Panel by panel", level=1)
    for p in brief.panels:
        d.add_paragraph(p.label, style="Heading 3")
        for i in p.items:
            d.add_paragraph(f"[{ACTION_LABEL.get(i.action, i.action)}] {i.item}: {i.instruction}", style="List Bullet")

    d.add_heading("3. Approved copy", level=1)
    table(["Content", "Copy", "Status"], [[c.label, c.value or "— not provided —", c.status.replace("_", " ")]
                                         for c in brief.final_copy], [1.6, 4.6, 1.0])
    d.add_heading("Nutrition Facts (per serving)", level=2)
    table(["Nutrient", "Amount", "%DV", "Status"],
          [[n["label"], n.get("amount") or "—", n.get("dv") or "", (n.get("status") or "").replace("_", " ")]
           for n in brief.nutrition], [2.4, 1.4, 1.0, 1.4])

    d.add_heading("4. Open questions", level=1)
    if brief.open_questions:
        table(["Question", "Owner", "Detail"], [[q.question, q.owner, q.detail[:400]] for q in brief.open_questions],
              [2.0, 0.9, 4.3])
    else:
        d.add_paragraph("None.")

    d.add_heading("5. Proof: where each value came from", level=1)
    table(["Field", "Value", "Status", "Sources"],
          [[p.label, (p.value or "—")[:160], p.status.replace("_", " "), "\n".join(p.sources)[:500]] for p in brief.proof],
          [1.3, 2.2, 0.8, 2.9])

    if brief.change_log:
        d.add_heading("6. Change log", level=1)
        table(["When (UTC)", "Who", "Action", "Detail"],
              [[c.at[:16].replace("T", " "), c.by, c.action,
                " ".join(x for x in [c.field or c.issue or "", f"'{c.old}' → '{c.new}'" if c.new is not None else "",
                                     c.note or ""] if x)[:300]] for c in brief.change_log], [1.2, 0.9, 0.9, 4.2])

    foot = d.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run("Generated by BriefForge. AI reads and suggests. Code checks. Humans decide.")
    fr.font.size = Pt(8)
    fr.font.italic = True
    d.save(str(path))
    return path


def export_pdf(brief: DesignerBrief, path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=8.5, leading=11)
    small = ParagraphStyle("s", parent=body, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#475569"))
    h1 = ParagraphStyle("h1", parent=ss["Heading2"], fontSize=13, spaceBefore=10, spaceAfter=4)

    def P(text, style=body):
        return Paragraph(escape(str(text or "")).replace("\n", "<br/>"), style)

    def tbl(headers, rows, widths):
        data = [[P(h, ParagraphStyle("hh", parent=body, textColor=colors.white, fontName="Helvetica-Bold")) for h in headers]]
        data += [[P(c) if not isinstance(c, Paragraph) else c for c in r] for r in rows]
        t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        return t

    story = [Paragraph("Packaging Design Brief", ss["Title"]),
             P(f"{brief.product} · {brief.market} · {brief.case}"),
             P(f"Status: {brief.status.upper()} · Updated {brief.updated_at[:16].replace('T', ' ')} UTC · "
               f"Readiness {brief.readiness.score if brief.readiness else '-'}%", small), Spacer(1, 4)]
    story += [P("• " + s) for s in brief.summary]

    story.append(Paragraph("1. What to do", h1))
    rows = []
    for i in brief.items:
        tag = Paragraph(f'<font color="#{ACTION_HEX.get(i.action, "334155")}"><b>{ACTION_LABEL.get(i.action, i.action)}</b></font>', body)
        rows.append([tag, i.item, _panel_name(i.panel), i.instruction,
                     (f"{i.current or ''} → {i.new or ''}" if i.action == "change" else (i.new or i.current or ""))[:220]])
    story.append(tbl(["Action", "Item", "Panel", "Instruction", "Current → New"], rows, [22, 31, 19, 55, 55]))

    story.append(Paragraph("2. Approved copy", h1))
    story.append(tbl(["Content", "Copy", "Status"],
                     [[c.label, c.value or "— not provided —", c.status.replace("_", " ")] for c in brief.final_copy],
                     [35, 125, 22]))
    story.append(Paragraph("Nutrition Facts (per serving)", h1))
    story.append(tbl(["Nutrient", "Amount", "%DV", "Status"],
                     [[n["label"], n.get("amount") or "—", n.get("dv") or "", (n.get("status") or "").replace("_", " ")]
                      for n in brief.nutrition], [60, 40, 30, 52]))

    story.append(Paragraph("3. Open questions", h1))
    story.append(tbl(["Question", "Owner", "Detail"],
                     [[q.question, q.owner, q.detail[:400]] for q in brief.open_questions] or [["None", "", ""]],
                     [50, 22, 110]))

    story.append(Paragraph("4. Proof: where each value came from", h1))
    story.append(tbl(["Field", "Value", "Status", "Sources"],
                     [[p.label, (p.value or "—")[:160], p.status.replace("_", " "), "\n".join(p.sources)[:500]]
                      for p in brief.proof], [30, 55, 18, 79]))
    if brief.change_log:
        story.append(Paragraph("5. Change log", h1))
        story.append(tbl(["When (UTC)", "Who", "Action", "Detail"],
                         [[c.at[:16].replace("T", " "), c.by, c.action,
                           " ".join(x for x in [c.field or c.issue or "", f"'{c.old}' → '{c.new}'" if c.new is not None else "",
                                                c.note or ""] if x)[:300]] for c in brief.change_log], [28, 22, 20, 112]))
    story.append(Spacer(1, 8))
    story.append(P("Generated by BriefForge. AI reads and suggests. Code checks. Humans decide.", small))
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm,
                      bottomMargin=14 * mm, title=f"Design brief - {brief.product}").build(story)
    return path
