"""Company rulebook: turn an uploaded document into proposed rules a reviewer can accept.

Every proposed rule names a brief field and quotes the sentence it came from. Code checks that the field exists and
that the quote really is in the document, so nothing can be invented. Rules found by plain keyword matching and rules
suggested by the AI go through the same checks. Nothing applies until a person accepts it.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Dict, List

from app.services.catalog import FIELD_BY_KEY, FIELDS
from app.services.guidelines import Clause, Requirement
from app.services.llm import get_llm
from app.services.values import norm_text

MAX_CHARS = 60000
MAX_RULES = 60
MUST = re.compile(r"\b(must|shall|required|mandatory|always|need to|needs to|has to|have to)\b", re.I)
SOFT = re.compile(r"\b(should|recommended|ensure|include)\b", re.I)


def extract_text(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext == "pdf":
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as doc:
            text = "\n".join(page.get_text() for page in doc)
    elif ext == "docx":
        from docx import Document

        d = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
            cell.text for t in d.tables for row in t.rows for cell in row.cells)
    elif ext in ("txt", "md"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError("Upload a PDF, DOCX, TXT or MD rulebook.")
    return re.sub(r"[ \t]+", " ", text).strip()[:MAX_CHARS]


def _sentences(text: str) -> List[str]:
    parts = re.split(r"(?<=[.;!?])\s+|\n+", text)
    return [p.strip(" -•*\t") for p in parts if 12 <= len(p.strip()) <= 400]


def keyword_rules(text: str) -> List[Dict]:
    """Sentences that say a thing 'must' / 'shall' appear and mention a known brief field."""
    out: List[Dict] = []
    for s in _sentences(text):
        strong, soft = bool(MUST.search(s)), bool(SOFT.search(s))
        if not (strong or soft):
            continue
        low = s.lower()
        for fd in FIELDS:
            names = {fd.label.lower(), *fd.synonyms}
            if any(len(n) >= 4 and re.search(rf"\b{re.escape(n)}\b", low) for n in names):
                out.append({"field": fd.key, "severity": "blocking" if strong else "warning",
                            "owner": fd.owner, "quote": s, "method": "keyword"})
    return out


def ai_rules(text: str) -> List[Dict]:
    model = get_llm("rules")
    if not model.enabled:
        return []
    fields = "\n".join(f"- {f.key}: {f.label}" for f in FIELDS)
    prompt = (
        "Below is a company's packaging rulebook. List rules that say a specific piece of information must appear "
        "on the pack (or in the brief). For each, pick the ONE field key from the list that it requires and copy the "
        "exact sentence from the rulebook as the quote. Never invent rules or change the wording of the quote.\n"
        'Return JSON only: {"rules":[{"field":"<key>","severity":"blocking|warning","quote":"<exact sentence>"}]}\n\n'
        f"FIELDS:\n{fields}\n\nRULEBOOK:\n{text[:30000]}")
    data = model.json(prompt)
    out = []
    for r in (data.get("rules") or []) if isinstance(data, dict) else []:
        if isinstance(r, dict) and r.get("field") in FIELD_BY_KEY and str(r.get("quote") or "").strip():
            out.append({"field": r["field"], "severity": "warning" if r.get("severity") == "warning" else "blocking",
                        "owner": FIELD_BY_KEY[r["field"]].owner, "quote": str(r["quote"]).strip(), "method": "ai"})
    return out


def propose(text: str) -> List[Dict]:
    """Verified proposals only: known field, and the quote really appears in the document."""
    doc = norm_text(text)
    seen, rules = set(), []
    try:
        candidates = keyword_rules(text) + ai_rules(text)
    except Exception:  # an AI failure must not lose the keyword rules
        candidates = keyword_rules(text)
    for c in candidates:
        key = (c["field"], norm_text(c["quote"]))
        if c["field"] not in FIELD_BY_KEY or key in seen or not norm_text(c["quote"]) or norm_text(c["quote"]) not in doc:
            continue
        seen.add(key)
        rules.append({**c, "label": FIELD_BY_KEY[c["field"]].label, "quote": c["quote"][:300]})
    rules = rules[:MAX_RULES]
    for i, r in enumerate(rules, start=1):
        r["id"] = f"r{i}"
    return rules


def to_clauses(accepted: List[Dict], filename: str) -> List[Clause]:
    """Accepted rules become clauses cited as COMPANY/<n>."""
    clauses = []
    for n, r in enumerate(accepted, start=1):
        field = r.get("field") if isinstance(r, dict) else None
        if field not in FIELD_BY_KEY or not str(r.get("quote") or "").strip():
            continue
        clauses.append(Clause(
            id=f"COMPANY/{n}", standard="company", standard_name="Company rulebook", section=FIELD_BY_KEY[field].section,
            text=str(r["quote"]).strip()[:300], keywords=[FIELD_BY_KEY[field].label.lower()],
            source={"name": f"Company rulebook: {filename}", "url": "", "note": "Accepted by the reviewer."},
            requires=[Requirement(field=field, severity="warning" if r.get("severity") == "warning" else "blocking",
                                  owner=r.get("owner") or FIELD_BY_KEY[field].owner)]))
    return clauses
