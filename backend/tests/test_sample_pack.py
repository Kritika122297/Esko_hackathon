"""Seeded-error test on the Esko sample pack, plus the human-in-the-loop API flow."""

import json
import shutil
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.config import settings
from app.main import app
from app.models.brief import BriefState
from app.services.pipeline import analyze, now, read_files, recompute

HERE = Path(__file__).parent
EXPECTED = json.loads((HERE / "expected_findings.json").read_text(encoding="utf-8"))


def _run(tmp_path, files):
    state = BriefState(session_id="test00000000", created_at=now(), updated_at=now())
    ex = read_files([(p, p.name) for p in files], tmp_path)
    ctx = analyze(state, ex)
    recompute(state, ctx)
    return state


@pytest.fixture(scope="module")
def sample_state(tmp_path_factory):
    files = sorted(settings.SAMPLES_DIR.glob("*"))
    assert files, "samples/ folder is empty"
    return _run(tmp_path_factory.mktemp("s"), files)


@pytest.mark.parametrize("case", EXPECTED["expected"] + EXPECTED["beyond_the_brief"], ids=lambda c: c["issue"])
def test_catches_seeded_error(sample_state, case):
    ids = {i.id for i in sample_state.issues}
    assert case["issue"] in ids, f"missed: {case['finding']}"


@pytest.mark.parametrize("case", EXPECTED["must_not_flag"], ids=lambda c: c["issue"])
def test_no_false_alarm(sample_state, case):
    ids = {i.id for i in sample_state.issues}
    assert case["issue"] not in ids, f"false alarm: {case['finding']}"


def test_values_have_proof(sample_state):
    for f in sample_state.fields:
        if f.value:
            assert f.sources, f"{f.key} has a value but no source"
        else:
            assert f.status == "missing"


def test_never_invents_values(sample_state):
    assert sample_state.value("registration") is None
    assert sample_state.value("last_edited") is None
    # The content sheet says the serving size was not legible: the only value is what the artwork prints.
    assert sample_state.value("serving_size") == "85"


def test_hidden_text_is_not_used_as_a_value(sample_state):
    for f in sample_state.fields:
        for c in f.candidates:
            assert "1868" not in c.value


def test_excel_only_run_still_works(tmp_path):
    xlsx = [p for p in settings.SAMPLES_DIR.glob("*.xlsx")]
    state = _run(tmp_path, xlsx)
    ids = {i.id for i in state.issues}
    assert "math.sugars_gt_carbs" in ids
    assert "us.net_weight_pdp" not in ids  # no artwork, no placement rule


def test_variation_of_the_data(tmp_path):
    """Changed content sheet: fixed sugars, a new net weight and a removed allergen line."""
    import openpyxl

    src = next(settings.SAMPLES_DIR.glob("*.xlsx"))
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    cells = {c.value: c for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    ws[cells["Total Sugars"].offset(column=1).coordinate] = "10g"
    ws[cells["Package Size / Net Weight"].offset(column=1).coordinate] = "500g"
    ws[cells["Allergen Information"].offset(column=1).coordinate] = None
    varied = tmp_path / "varied.xlsx"
    wb.save(varied)
    pdf = next(settings.SAMPLES_DIR.glob("*.pdf"))
    state = _run(tmp_path, [pdf, varied])
    ids = {i.id for i in state.issues}
    assert "math.sugars_gt_carbs" not in ids          # 10g <= 15g now
    assert "conflict:total_sugars" in ids             # artwork still prints 18g
    assert "conflict:net_weight" in ids               # 500g vs 480g on pack
    # The allergen statement now comes only from the artwork, and is still checked
    assert state.field("allergen_statement").sources[0].kind == "pdf"
    assert "allergen.may_contain_present" in ids


@pytest.fixture()
def client(tmp_path, monkeypatch):
    from app import session_store

    monkeypatch.setattr(session_store.store, "root", tmp_path)
    monkeypatch.setattr(settings, "UPLOAD_DIR", tmp_path)
    return TestClient(app)


def test_api_flow(client):
    # The sample pack is uploaded like any other product's files.
    files = [("files", (p.name, p.read_bytes())) for p in sorted(settings.SAMPLES_DIR.glob("*"))]
    r = client.post("/api/sessions", files=files, data={"case": "update", "market": "US"})
    assert r.status_code == 201, r.text
    view = r.json()
    sid = view["state"]["session_id"]
    assert view["pages"] and view["panels"]
    assert client.get(view["pages"][0]["url"]).status_code == 200

    # Finalize is blocked while blocking decisions are open
    r = client.post(f"/api/sessions/{sid}/finalize")
    assert r.status_code == 409

    # Answer the net-weight question with the code-computed option
    q = next(q for q in view["state"]["questions"] if "us.net_weight_oz" in q["issue_ids"])
    opt = next(o for o in q["options"] if o["updates"])
    r = client.post(f"/api/sessions/{sid}/questions/{q['id']}/answer", json={"option_id": opt["id"]})
    state = r.json()["state"]
    assert "us.net_weight_oz" not in {i["id"] for i in state["issues"]}
    assert next(f for f in state["fields"] if f["key"] == "net_weight")["value"].startswith("NET WT 16.9 OZ")

    # Impact preview does not change anything
    r = client.post(f"/api/sessions/{sid}/impact", json={"field": "total_carbohydrate", "value": "66g"})
    assert r.status_code == 200
    preview = r.json()
    assert "math.sugars_gt_carbs" in {i["id"] for i in preview["resolved"]}
    carbs = next(f for f in client.get(f"/api/sessions/{sid}").json()["state"]["fields"] if f["key"] == "total_carbohydrate")
    assert carbs["value"] == "15g"

    # Acknowledge every remaining blocking data issue -> finalize allowed
    state = client.get(f"/api/sessions/{sid}").json()["state"]
    for i in state["issues"]:
        if i["state"] == "open" and i["severity"] == "blocking" and i["resolution"] == "data":
            client.post(f"/api/sessions/{sid}/issues/{i['id']}/acknowledge", json={"note": "Agency to confirm"})
    r = client.post(f"/api/sessions/{sid}/finalize")
    assert r.status_code == 200, r.text
    assert r.json()["state"]["status"] == "finalized"

    brief = client.get(f"/api/sessions/{sid}/designer-brief").json()
    assert {i["action"] for i in brief["items"]} >= {"keep", "change", "remove"}
    assert brief["open_questions"]
    for fmt in ("docx", "pdf"):
        r = client.get(f"/api/sessions/{sid}/export/{fmt}")
        assert r.status_code == 200 and len(r.content) > 2000


def test_projects_list_chat_and_delete(client):
    files = [("files", (p.name, p.read_bytes())) for p in sorted(settings.SAMPLES_DIR.glob("*"))]
    sid = client.post("/api/sessions", files=files, data={"case": "update", "market": "US"}).json()["state"]["session_id"]

    projects = client.get("/api/sessions").json()
    assert [p["session_id"] for p in projects] == [sid]
    assert projects[0]["blocking_open"] > 0

    a = client.post(f"/api/sessions/{sid}/chat", json={"message": "What is still blocking the brief?"}).json()
    assert "blocking" in a["answer"].lower() and a["citations"]
    a = client.post(f"/api/sessions/{sid}/chat", json={"message": "what is the total carbohydrate?"}).json()
    assert "15g" in a["answer"] and "66g" in a["answer"]
    assert a["citations"][0]["key"] == "total_carbohydrate"

    assert client.delete(f"/api/sessions/{sid}").status_code == 204
    assert client.get("/api/sessions").json() == []
    assert client.get(f"/api/sessions/{sid}").status_code == 404


def test_duplicate_upload_add_files_and_chat_history(client):
    pdf = next(settings.SAMPLES_DIR.glob("*.pdf"))
    xlsx = next(settings.SAMPLES_DIR.glob("*.xlsx"))
    up = lambda paths, **extra: client.post(  # noqa: E731
        "/api/sessions", files=[("files", (p.name, p.read_bytes())) for p in paths], data={"case": "update", "market": "US", **extra})

    # Start from the content sheet only, then add the artwork to the same brief.
    first = up([xlsx])
    sid = first.json()["state"]["session_id"]
    tc = next(f for f in first.json()["state"]["fields"] if f["key"] == "total_carbohydrate")
    assert all(c["source"]["kind"] != "pdf" for c in tc["candidates"])

    # A reviewer decision made before the new file arrives must survive the re-read.
    client.put(f"/api/sessions/{sid}/fields/brand", json={"value": "Dancing Fish", "reason": "checked"})
    r = client.post(f"/api/sessions/{sid}/files", files=[("files", (pdf.name, pdf.read_bytes()))])
    assert r.status_code == 200, r.text
    st = r.json()["state"]
    assert {f["name"] for f in st["files"]} == {pdf.name, xlsx.name}
    assert "conflict:total_carbohydrate" in {i["id"] for i in st["issues"]}  # artwork now compared
    brand = next(f for f in st["fields"] if f["key"] == "brand")
    assert brand["user_confirmed"] and brand["history"]
    assert st["change_log"][-1]["action"] == "files"
    assert len(client.get("/api/sessions").json()) == 1

    # Uploading the same files again points to the existing brief instead of duplicating it.
    dup = up([pdf, xlsx])
    assert dup.status_code == 409 and dup.json()["detail"]["existing"] == sid
    assert up([pdf, xlsx], allow_duplicate="true").status_code == 201

    # Chat history is saved with the brief.
    client.post(f"/api/sessions/{sid}/chat", json={"message": "Summarize this brief"})
    hist = client.get(f"/api/sessions/{sid}/chat").json()
    assert [m["role"] for m in hist] == ["user", "assistant"]
    assert client.delete(f"/api/sessions/{sid}/chat").status_code == 204
    assert client.get(f"/api/sessions/{sid}/chat").json() == []
