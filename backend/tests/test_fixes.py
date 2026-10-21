"""Regression tests for the review findings: each test names the problem it guards against."""

import io

from app.config import settings
from test_sample_pack import client  # noqa: F401  (isolated session store)
from test_sample_pack_extra import _session


def _files():
    return [("files", (p.name, p.read_bytes())) for p in sorted(settings.SAMPLES_DIR.glob("*")) if p.is_file()]


def _state(c, sid):
    return c.get(f"/api/sessions/{sid}").json()["state"]


# ---------------------------------------------------------------- checks


def test_plant_ingredients_are_not_animal_products_or_milk_allergens():
    from app.services.checks.claims import ANIMAL_DERIVED, _allergens_in, _matches

    lists = [("Coconut milk, eggplant, cocoa butter, shea butter, buckwheat", None)]
    assert _matches(ANIMAL_DERIVED, lists) == {}
    assert _allergens_in("Coconut milk, eggplant, cocoa butter, buckwheat") == []
    assert set(_matches(ANIMAL_DERIVED, [("Honey, sodium caseinate, gelatine, eggs", None)])) == {"honey", "casein", "gelatin", "egg"}
    assert set(_allergens_in("Almond milk, whole milk, soya lecithin")) == {"Almonds", "Milk", "Soy"}


def test_removing_the_vegan_claim_resolves_it_and_tells_the_designer(client):
    c = client
    sid = _session(c)
    q = next(q for q in _state(c, sid)["questions"] if "claim.vegan_animal" in q["issue_ids"])
    opt = next(o for o in q["options"] if o["recommended"])
    st = c.post(f"/api/sessions/{sid}/questions/{q['id']}/answer", json={"option_id": opt["id"]}).json()["state"]
    ids = {i["id"] for i in st["issues"]}
    assert "claim.vegan_animal" not in ids and "claim.vegan_cholesterol" not in ids
    assert q["id"] not in {x["id"] for x in st["questions"]}  # not asked again
    brief = c.get(f"/api/sessions/{sid}/designer-brief").json()
    assert any(i["action"] == "remove" and "vegan" in i["item"].lower() for i in brief["items"])


def test_amounts_read_decimal_commas_and_pounds_with_ounces():
    from app.services.values import amount_in, parse_amount

    assert parse_amount("3,5 g") == (3.5, "g")
    assert parse_amount("1,000 g") == (1000.0, "g")
    assert amount_in("NET WT 1 LB 4 OZ (567g)", "g") == 567.0
    assert amount_in("NET WT 1 LB 4 OZ", "oz") == 20.0
    assert amount_in("12 fl oz", "g") is None  # volume is not weight


def test_warning_requirements_do_not_block(client):
    from app.services.checks import run_checks
    from app.services.guidelines import requirements
    from app.session_store import store

    sid = _session(client)
    state, ctx, _ = store.load(sid)
    st = state.model_copy(deep=True)
    st.category = "medical_device"
    reqs = requirements(st.market, st.category)
    warn_only = [k for k, pairs in reqs.items() if all(r.severity == "warning" for _, r in pairs)]
    assert warn_only
    for f in st.fields:
        f.mandatory = f.key in reqs
        if f.key in warn_only:
            f.value, f.not_applicable = None, False
    issues, _ = run_checks(st, ctx)
    assert all(i.severity == "warning" for i in issues if i.check == "missing" and i.fields[0] in warn_only)


def test_theme_coloured_excel_cells_do_not_break_the_sheet():
    from openpyxl import Workbook
    from openpyxl.styles import Color, PatternFill

    from app.services.file_parser import _is_highlight

    ws = Workbook().active
    ws["A1"] = "x"
    ws["A1"].fill = PatternFill(fill_type="solid", fgColor=Color(theme=4))
    assert _is_highlight(ws["A1"]) is False


# ---------------------------------------------------------------- review actions


def test_undo_of_not_applicable_makes_the_field_apply_again(client):
    c = client
    sid = _session(c)
    c.put(f"/api/sessions/{sid}/fields/added_sugars", json={"value": ""})
    c.post(f"/api/sessions/{sid}/fields/added_sugars/not-applicable", json={"reason": "Unsweetened"})
    st = c.post(f"/api/sessions/{sid}/fields/added_sugars/undo").json()["state"]
    f = next(f for f in st["fields"] if f["key"] == "added_sugars")
    assert not f["not_applicable"] and f["absence"] != "not_applicable"
    st = c.post(f"/api/sessions/{sid}/fields/added_sugars/undo").json()["state"]  # undo again re-applies it
    assert next(f for f in st["fields"] if f["key"] == "added_sugars")["not_applicable"]


def test_undo_back_to_a_reviewer_value_keeps_it_confirmed(client):
    c = client
    sid = _session(c)
    c.put(f"/api/sessions/{sid}/fields/storage", json={"value": "Keep dry"})
    c.put(f"/api/sessions/{sid}/fields/storage", json={"value": "Keep cool"})
    st = c.post(f"/api/sessions/{sid}/fields/storage/undo").json()["state"]
    f = next(f for f in st["fields"] if f["key"] == "storage")
    assert f["value"] == "Keep dry" and f["user_confirmed"]


def test_word_export_survives_control_characters(client):
    c = client
    sid = _session(c)
    c.put(f"/api/sessions/{sid}/fields/storage", json={"value": "Keep dry\x0bStore cool\x01"})
    assert c.get(f"/api/sessions/{sid}/export/docx").status_code == 200


# ---------------------------------------------------------------- uploads


def test_duplicate_file_names_are_refused_before_anything_is_saved(client):
    from app.session_store import store

    before = set(store.list_ids())
    files = _files()
    r = client.post("/api/sessions", files=files + [files[0]], data={"case": "update", "market": "US"})
    assert r.status_code == 400 and "named" in r.json()["detail"]
    assert set(store.list_ids()) == before


def test_a_rejected_batch_does_not_replace_existing_files(client):
    from app.session_store import store

    sid = _session(client)
    sheet = next(f for f in store.load(sid)[0].files if f.name.endswith(".xlsx"))
    original = (store.dir(sid) / sheet.name).read_bytes()
    bad = [("files", (sheet.name, b"not a spreadsheet")), ("files", ("x.exe", b"MZ"))]
    assert client.post(f"/api/sessions/{sid}/files", files=bad).status_code == 400
    assert (store.dir(sid) / sheet.name).read_bytes() == original


def test_a_finalized_brief_takes_no_files_until_reopened(client):
    from app.session_store import store

    sid = _session(client)
    store.load(sid)[0].status = "finalized"
    r = client.post(f"/api/sessions/{sid}/files", files=[("files", ("notes.txt", b"Storage: keep dry"))])
    assert r.status_code == 409
    assert _state(client, sid)["status"] == "finalized"


def test_one_damaged_brief_does_not_break_the_project_list(client):
    from app.session_store import store

    good = _session(client)
    bad = client.post("/api/sessions", files=_files(),
                      data={"case": "update", "market": "US", "allow_duplicate": "true"}).json()["state"]["session_id"]
    store._cache.pop(bad)
    (store.dir(bad) / "state.json").write_text("{ not json", encoding="utf-8")
    listed = client.get("/api/sessions")
    assert listed.status_code == 200 and good in {p["session_id"] for p in listed.json()}
    assert client.get(f"/api/sessions/{bad}").status_code == 422
    assert client.delete(f"/api/sessions/{bad}").status_code == 204


def test_file_names_with_hash_signs_get_working_page_links(client):
    pdf = next(p for p in settings.SAMPLES_DIR.glob("*.pdf"))
    r = client.post("/api/sessions", files=[("files", ("Carton #2.pdf", pdf.read_bytes()))],
                    data={"case": "update", "market": "US", "allow_duplicate": "true"})
    url = r.json()["pages"][0]["url"]
    assert "#" not in url and client.get(url).status_code == 200


# ---------------------------------------------------------------- the three situations


def test_new_market_uses_the_new_spec_and_asks_about_carried_over_copy(client):
    r = client.post("/api/sessions", files=_files(),
                    data={"case": "new_market", "market": "US", "source_market": "EU", "allow_duplicate": "true"})
    st = r.json()["state"]
    assert st["source_market"] == "EU"
    # Artwork (old pack) against content sheet (new market) is the change to make, not a conflict question:
    # the same files as a pack refresh do raise these conflicts.
    refresh = {i["id"] for i in client.post("/api/sessions", files=_files(), data={
        "case": "update", "market": "US", "allow_duplicate": "true"}).json()["state"]["issues"]}
    assert {"conflict:sodium", "conflict:total_carbohydrate"} <= refresh
    assert not {"conflict:sodium", "conflict:total_carbohydrate"} & {i["id"] for i in st["issues"]}
    brief = client.get(f"/api/sessions/{st['session_id']}/designer-brief").json()
    assert brief["summary"][0].startswith("Adapt the EU pack for the US market")


def test_carried_over_values_are_asked_about_for_the_new_market(client):
    from app.services.checks import run_checks
    from app.session_store import store

    sid = _session(client)
    state, ctx, _ = store.load(sid)
    st = state.model_copy(deep=True)
    st.case, st.source_market = "new_market", "EU"
    f = st.field("producer")
    f.sources = [s for s in f.sources if s.kind == "pdf"] or f.sources
    for s in f.sources:
        s.kind = "pdf"
    f.mandatory, f.user_confirmed = True, False
    issues, _ = run_checks(st, ctx)
    covered = {k for i in issues if i.check != "carryover" and i.severity in ("blocking", "warning") for k in i.fields}
    if "producer" not in covered:
        assert any(i.id == "carryover:producer" for i in issues)


def test_new_product_treats_artwork_as_reference_only(client):
    r = client.post("/api/sessions", files=_files(), data={"case": "new_product", "market": "US", "allow_duplicate": "true"})
    st = r.json()["state"]
    fields = {f["key"]: f for f in st["fields"]}
    # Values only the artwork had are not used.
    assert all(not f["value"] or any(s["kind"] != "pdf" for s in f["sources"]) for f in fields.values())
    assert not any(i["check"] in ("artwork.hidden_text", "us.net_weight_pdp", "conflict") and
                   {e["kind"] for e in i["evidence"]} >= {"pdf"} and len({e["kind"] for e in i["evidence"]}) == 1
                   for i in st["issues"])
    brief = client.get(f"/api/sessions/{st['session_id']}/designer-brief").json()
    assert brief["summary"][0].startswith("New product")
    assert not any(i["action"] in ("keep", "change") for i in brief["items"])
